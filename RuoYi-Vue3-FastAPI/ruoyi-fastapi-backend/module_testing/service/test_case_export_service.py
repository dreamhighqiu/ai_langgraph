"""
测试用例导出服务
支持导出为Excel和XMind格式
"""
import json
from typing import List, Optional
from datetime import datetime
from io import BytesIO
from zipfile import ZipFile
import xml.etree.ElementTree as ET
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from module_testing.entity.do.test_case_do import TestCaseDO
from utils.log_util import logger


class TestCaseExportService:
    """测试用例导出服务"""
    
    def __init__(self):
        pass
    
    def export_to_excel(self, test_cases: List[TestCaseDO], project_name: str = "测试项目") -> BytesIO:
        """
        导出测试用例到Excel
        
        Args:
            test_cases: 测试用例列表
            project_name: 项目名称
        
        Returns:
            BytesIO: Excel文件流
        """
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "测试用例"
            
            # 设置列宽
            column_widths = {
                'A': 12,  # 用例编号
                'B': 40,  # 用例名称
                'C': 15,  # 用例类型
                'D': 10,  # 优先级
                'E': 10,  # 严重程度
                'F': 15,  # 状态
                'G': 15,  # 执行状态
                'H': 30,  # 前置条件
                'I': 50,  # 测试步骤
                'J': 50,  # 预期结果
                'K': 30,  # 测试数据
                'L': 30,  # 标签
                'M': 20,  # 创建人
                'N': 20,  # 创建时间
            }
            
            for col, width in column_widths.items():
                ws.column_dimensions[col].width = width
            
            # 定义样式
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
            header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            
            border_style = Border(
                left=Side(style='thin', color='000000'),
                right=Side(style='thin', color='000000'),
                top=Side(style='thin', color='000000'),
                bottom=Side(style='thin', color='000000')
            )
            
            content_font = Font(name="微软雅黑", size=10)
            content_alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            
            # 创建标题行
            title_row = 1
            ws.merge_cells(f'A{title_row}:N{title_row}')
            title_cell = ws[f'A{title_row}']
            title_cell.value = f"{project_name} - 测试用例清单"
            title_cell.font = Font(name="微软雅黑", size=16, bold=True)
            title_cell.alignment = Alignment(horizontal="center", vertical="center")
            title_cell.fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
            ws.row_dimensions[title_row].height = 30
            
            # 创建表头
            headers = [
                "用例编号", "用例名称", "用例类型", "优先级", "严重程度", 
                "状态", "执行状态", "前置条件", "测试步骤", "预期结果", 
                "测试数据", "标签", "创建人", "创建时间"
            ]
            
            header_row = 3
            for col_idx, header in enumerate(headers, start=1):
                cell = ws.cell(row=header_row, column=col_idx)
                cell.value = header
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = border_style
            
            ws.row_dimensions[header_row].height = 25
            
            # 填充数据
            for row_idx, test_case in enumerate(test_cases, start=header_row + 1):
                # 解析测试步骤和预期结果，分别存储
                test_steps = []
                expected_results = []
                
                if test_case.test_case_steps:
                    try:
                        steps_data = json.loads(test_case.test_case_steps) if isinstance(test_case.test_case_steps, str) else test_case.test_case_steps
                        if isinstance(steps_data, list):
                            for idx, step in enumerate(steps_data, start=1):
                                if isinstance(step, dict):
                                    # 支持多种字段名格式
                                    action = step.get('action') or step.get('step_description') or step.get('step_action') or ''
                                    expected = step.get('expected') or step.get('expected_result') or step.get('step_expected') or ''
                                    
                                    if action or expected:
                                        # 步骤和预期结果分开存储
                                        test_steps.append(f"{idx}. {action}" if action else f"{idx}. ")
                                        expected_results.append(f"{idx}. {expected}" if expected else f"{idx}. ")
                                elif isinstance(step, str):
                                    # 如果是字符串，直接使用
                                    test_steps.append(f"{idx}. {step}")
                                    expected_results.append(f"{idx}. ")
                    except Exception as e:
                        logger.warning(f"解析测试步骤失败: {str(e)}, 原始数据: {test_case.test_case_steps}")
                
                steps_text = "\n".join(test_steps) if test_steps else ""
                expected_text = "\n".join(expected_results) if expected_results else ""
                
                # 解析标签
                tags = []
                if test_case.tags:
                    tags = test_case.tags.split(',') if isinstance(test_case.tags, str) else test_case.tags
                tags_text = ", ".join(tags) if tags else ""
                
                # 格式化优先级
                priority_map = {
                    'critical': '关键',
                    'high': '高',
                    'medium': '中',
                    'low': '低'
                }
                priority_text = priority_map.get(test_case.priority, test_case.priority or '中')
                
                # 格式化状态
                status_map = {
                    'draft': '草稿',
                    'active': '激活',
                    'deprecated': '废弃',
                    'pending': '待评审',
                    'approved': '已通过'
                }
                status_text = status_map.get(test_case.status, test_case.status or '草稿')
                
                # 格式化执行状态
                exec_status_map = {
                    'not_run': '未执行',
                    'passed': '通过',
                    'failed': '失败',
                    'blocked': '阻塞'
                }
                exec_status_text = exec_status_map.get(test_case.execution_status, test_case.execution_status or '未执行')
                
                # 格式化用例类型
                case_type_map = {
                    'functional': '功能测试',
                    'regression': '回归测试',
                    'smoke': '冒烟测试',
                    'performance': '性能测试',
                    'security': '安全测试'
                }
                case_type_text = case_type_map.get(test_case.case_type, test_case.case_type or '功能测试')
                
                # 格式化时间
                create_time_text = ""
                if test_case.create_time:
                    if isinstance(test_case.create_time, str):
                        create_time_text = test_case.create_time
                    else:
                        create_time_text = test_case.create_time.strftime('%Y-%m-%d %H:%M:%S')
                
                # 填充数据
                data = [
                    test_case.case_identifier or "",
                    test_case.case_name or "",
                    case_type_text,
                    priority_text,
                    test_case.severity or "",
                    status_text,
                    exec_status_text,
                    test_case.preconditions or "",
                    steps_text,  # 测试步骤
                    expected_text,  # 预期结果（与步骤一一对应）
                    test_case.test_data or "",
                    tags_text,
                    test_case.create_by or "",
                    create_time_text
                ]
                
                for col_idx, value in enumerate(data, start=1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.value = value
                    cell.font = content_font
                    cell.alignment = content_alignment
                    cell.border = border_style
                    
                    # 设置行高（根据内容自动调整）
                    if col_idx in [2, 8, 9, 10, 11]:  # 长文本列
                        ws.row_dimensions[row_idx].height = max(30, len(str(value).split('\n')) * 15)
            
            # 冻结表头
            ws.freeze_panes = f'A{header_row + 1}'
            
            # 创建表格（可选，增强视觉效果）
            if len(test_cases) > 0:
                table = Table(displayName="TestCaseTable", ref=f"A{header_row}:N{header_row + len(test_cases)}")
                style = TableStyleInfo(
                    name="TableStyleMedium9",
                    showFirstColumn=False,
                    showLastColumn=False,
                    showRowStripes=True,
                    showColumnStripes=False
                )
                table.tableStyleInfo = style
                ws.add_table(table)
            
            # 保存到内存
            output = BytesIO()
            wb.save(output)
            output.seek(0)
            
            logger.info(f"成功导出 {len(test_cases)} 条测试用例到Excel")
            return output
            
        except Exception as e:
            logger.error(f"导出Excel失败: {str(e)}", exc_info=True)
            raise
    
    def export_to_xmind(self, test_cases: List[TestCaseDO], project_name: str = "测试项目") -> BytesIO:
        """
        导出测试用例到XMind
        
        Args:
            test_cases: 测试用例列表
            project_name: 项目名称
        
        Returns:
            BytesIO: XMind文件流
        """
        try:
            # XMind文件实际上是ZIP格式，包含多个XML文件
            
            # 创建XMind内容结构
            # XMind文件结构：
            # - content.xml (主要内容)
            # - meta.xml (元数据)
            # - META-INF/manifest.xml (清单文件)
            
            output = BytesIO()
            
            with ZipFile(output, 'w') as zip_file:
                # 创建content.xml
                content_xml = self._create_xmind_content(test_cases, project_name)
                zip_file.writestr('content.xml', content_xml.encode('utf-8'))
                
                # 创建meta.xml
                meta_xml = self._create_xmind_meta(project_name)
                zip_file.writestr('meta.xml', meta_xml.encode('utf-8'))
                
                # 创建META-INF/manifest.xml
                manifest_xml = self._create_xmind_manifest()
                zip_file.writestr('META-INF/manifest.xml', manifest_xml.encode('utf-8'))
            
            output.seek(0)
            
            logger.info(f"成功导出 {len(test_cases)} 条测试用例到XMind")
            return output
            
        except Exception as e:
            logger.error(f"导出XMind失败: {str(e)}", exc_info=True)
            raise
    
    def _create_xmind_content(self, test_cases: List[TestCaseDO], project_name: str) -> str:
        """创建XMind content.xml内容"""
        
        # 创建根元素
        root = ET.Element('xmap-content', {
            'xmlns': 'urn:xmind:xmap:xmlns:content:2.0',
            'xmlns:fo': 'http://www.w3.org/1999/XSL/Format',
            'version': '2.0'
        })
        
        # 创建工作表
        sheet = ET.SubElement(root, 'sheet', {'id': 'root'})
        
        # 创建主题（根节点）
        topic = ET.SubElement(sheet, 'topic', {'id': 'root-topic'})
        title = ET.SubElement(topic, 'title')
        title.text = project_name
        
        # 按用例类型分组
        type_groups = {}
        for test_case in test_cases:
            case_type = test_case.case_type or '功能测试'
            if case_type not in type_groups:
                type_groups[case_type] = []
            type_groups[case_type].append(test_case)
        
        # 创建子主题（用例类型）
        for case_type, cases in type_groups.items():
            type_topic = ET.SubElement(topic, 'children')
            type_topics = ET.SubElement(type_topic, 'topics', {'type': 'attached'})
            type_topic_elem = ET.SubElement(type_topics, 'topic', {'id': f'type-{case_type}'})
            type_title = ET.SubElement(type_topic_elem, 'title')
            type_title.text = case_type
            
            # 创建用例子主题
            cases_topics = ET.SubElement(type_topic_elem, 'children')
            cases_topics_elem = ET.SubElement(cases_topics, 'topics', {'type': 'attached'})
            
            for test_case in cases:
                case_topic = ET.SubElement(cases_topics_elem, 'topic', {'id': f'case-{test_case.case_id}'})
                case_title = ET.SubElement(case_topic, 'title')
                case_title.text = f"{test_case.case_identifier} - {test_case.case_name}"
                
                # 添加用例详细信息作为备注
                notes = ET.SubElement(case_topic, 'notes')
                notes_content = ET.SubElement(notes, 'plain', {'xmlns': 'http://www.w3.org/1999/xhtml'})
                
                note_text = []
                if test_case.description:
                    note_text.append(f"描述: {test_case.description}")
                if test_case.preconditions:
                    note_text.append(f"前置条件: {test_case.preconditions}")
                if test_case.test_data:
                    note_text.append(f"测试数据: {test_case.test_data}")
                
                notes_content.text = "\n\n".join(note_text) if note_text else "无"
                
                # 添加测试步骤作为子主题（步骤和预期结果一一对应）
                if test_case.test_case_steps:
                    try:
                        steps_data = json.loads(test_case.test_case_steps) if isinstance(test_case.test_case_steps, str) else test_case.test_case_steps
                        if isinstance(steps_data, list) and steps_data:
                            # 创建测试步骤父节点
                            steps_children = ET.SubElement(case_topic, 'children')
                            steps_topics = ET.SubElement(steps_children, 'topics', {'type': 'attached'})
                            
                            for idx, step in enumerate(steps_data, start=1):
                                if isinstance(step, dict):
                                    # 支持多种字段名格式
                                    action = step.get('action') or step.get('step_description') or step.get('step_action') or ''
                                    expected = step.get('expected') or step.get('expected_result') or step.get('step_expected') or ''
                                    
                                    if action or expected:
                                        # 创建步骤节点
                                        step_topic = ET.SubElement(steps_topics, 'topic', {'id': f'step-{test_case.case_id}-{idx}'})
                                        step_title = ET.SubElement(step_topic, 'title')
                                        step_title.text = f"步骤{idx}: {action if action else '(无)'}"
                                        
                                        # 如果有预期结果，添加为子节点
                                        if expected:
                                            expected_children = ET.SubElement(step_topic, 'children')
                                            expected_topics = ET.SubElement(expected_children, 'topics', {'type': 'attached'})
                                            expected_topic = ET.SubElement(expected_topics, 'topic', {'id': f'expected-{test_case.case_id}-{idx}'})
                                            expected_title = ET.SubElement(expected_topic, 'title')
                                            expected_title.text = f"预期结果: {expected}"
                    except Exception as e:
                        logger.warning(f"解析XMind测试步骤失败: {str(e)}, 原始数据: {test_case.test_case_steps}")
                
                # 添加标签
                if test_case.tags:
                    labels = ET.SubElement(case_topic, 'labels')
                    tags = test_case.tags.split(',') if isinstance(test_case.tags, str) else test_case.tags
                    for tag in tags:
                        if tag.strip():
                            label = ET.SubElement(labels, 'label')
                            label.text = tag.strip()
        
        # 转换为XML字符串
        ET.indent(root, space="  ")
        return ET.tostring(root, encoding='unicode', xml_declaration=True)
    
    def _create_xmind_meta(self, project_name: str) -> str:
        """创建XMind meta.xml内容"""
        
        root = ET.Element('meta', {'xmlns': 'urn:xmind:xmap:xmlns:meta:2.0'})
        author = ET.SubElement(root, 'Author')
        author.text = 'Test Case Export Service'
        
        ET.indent(root, space="  ")
        return ET.tostring(root, encoding='unicode', xml_declaration=True)
    
    def _create_xmind_manifest(self) -> str:
        """创建XMind manifest.xml内容"""
        
        root = ET.Element('manifest', {'xmlns': 'urn:xmind:xmap:xmlns:manifest:1.0'})
        
        # 添加文件条目
        file_entry = ET.SubElement(root, 'file-entry', {
            'full-path': 'content.xml',
            'media-type': 'text/xml'
        })
        
        file_entry2 = ET.SubElement(root, 'file-entry', {
            'full-path': 'meta.xml',
            'media-type': 'text/xml'
        })
        
        ET.indent(root, space="  ")
        return ET.tostring(root, encoding='unicode', xml_declaration=True)

