"""
Excel 导出工具

将测试用例导出为美观的 Excel 格式
"""

from typing import List, Dict, Any
from pathlib import Path
from datetime import datetime

try:
    import openpyxl
    from openpyxl.styles import (
        Font, PatternFill, Alignment, Border, Side,
        GradientFill, NamedStyle
    )
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.formatting.rule import FormulaRule, CellIsRule
    from openpyxl.comments import Comment
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

from .testcase_tools import TestCase, Priority, TestType


class ExcelExporter:
    """Excel 导出器 - 生成美观的测试用例文档"""
    
    # ==================== 颜色主题 ====================
    # 深色专业主题
    THEME = {
        'primary': '2F5496',       # 深蓝色 - 主色调
        'secondary': '4472C4',     # 中蓝色
        'accent': '5B9BD5',        # 亮蓝色
        'header_text': 'FFFFFF',   # 白色文字
        'dark_text': '1F2937',     # 深灰文字
        'light_text': '6B7280',    # 浅灰文字
        'row_even': 'F8FAFC',      # 偶数行背景 - 极浅灰
        'row_odd': 'FFFFFF',       # 奇数行背景 - 白色
        'border': 'D1D5DB',        # 边框色 - 浅灰
        'border_header': '2F5496', # 表头边框
    }
    
    # 优先级颜色 - 使用更柔和的颜色
    PRIORITY_COLORS = {
        "P0": {
            'bg': 'FEE2E2',    # 浅红背景
            'text': 'DC2626',  # 红色文字
            'border': 'FECACA'
        },
        "P1": {
            'bg': 'FEF3C7',    # 浅橙背景
            'text': 'D97706',  # 橙色文字
            'border': 'FDE68A'
        },
        "P2": {
            'bg': 'D1FAE5',    # 浅绿背景
            'text': '059669',  # 绿色文字
            'border': 'A7F3D0'
        },
        "P3": {
            'bg': 'E0E7FF',    # 浅紫背景
            'text': '4F46E5',  # 紫色文字
            'border': 'C7D2FE'
        },
    }
    
    # 测试类型颜色
    TYPE_COLORS = {
        "功能测试": {'bg': 'DBEAFE', 'text': '1D4ED8'},
        "functional": {'bg': 'DBEAFE', 'text': '1D4ED8'},
        "UI测试": {'bg': 'E0E7FF', 'text': '4F46E5'},
        "ui": {'bg': 'E0E7FF', 'text': '4F46E5'},
        "冒烟测试": {'bg': 'FEF3C7', 'text': 'B45309'},
        "smoke": {'bg': 'FEF3C7', 'text': 'B45309'},
        "边界测试": {'bg': 'ECFDF5', 'text': '047857'},
        "boundary": {'bg': 'ECFDF5', 'text': '047857'},
        "异常测试": {'bg': 'FEE2E2', 'text': 'B91C1C'},
        "negative": {'bg': 'FEE2E2', 'text': 'B91C1C'},
    }
    
    # ==================== 列配置 ====================
    # UI测试用例列定义 - 包含定位器
    UI_COLUMNS = [
        {"name": "测试用例ID", "width": 14, "key": "id"},
        {"name": "模块", "width": 12, "key": "module"},
        {"name": "子模块", "width": 12, "key": "submodule"},
        {"name": "测试场景", "width": 25, "key": "title"},
        {"name": "测试步骤", "width": 55, "key": "steps"},
        {"name": "预期结果", "width": 25, "key": "expected_result"},
        {"name": "优先级", "width": 8, "key": "priority"},
        {"name": "测试类型", "width": 12, "key": "type"},
        {"name": "测试数据", "width": 20, "key": "test_data"},
        {"name": "备注", "width": 15, "key": "notes"},
    ]
    
    # 全量测试用例列定义
    ALL_COLUMNS = [
        {"name": "测试用例ID", "width": 14, "key": "id"},
        {"name": "模块", "width": 12, "key": "module"},
        {"name": "子模块", "width": 12, "key": "submodule"},
        {"name": "测试场景", "width": 25, "key": "title"},
        {"name": "测试步骤", "width": 55, "key": "steps"},
        {"name": "预期结果", "width": 25, "key": "expected_result"},
        {"name": "优先级", "width": 8, "key": "priority"},
        {"name": "测试类型", "width": 12, "key": "type"},
        {"name": "测试数据", "width": 20, "key": "test_data"},
        {"name": "备注", "width": 15, "key": "notes"},
    ]
    
    def __init__(self):
        if not OPENPYXL_AVAILABLE:
            raise ImportError("openpyxl 未安装。请运行: pip install openpyxl")
        
        self._init_styles()
    
    def _init_styles(self):
        """初始化样式"""
        # 表头样式
        self.header_font = Font(
            name='微软雅黑',
            bold=True,
            color=self.THEME['header_text'],
            size=11
        )
        self.header_fill = PatternFill(
            start_color=self.THEME['primary'],
            end_color=self.THEME['primary'],
            fill_type="solid"
        )
        self.header_alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )
        self.header_border = Border(
            left=Side(style='thin', color=self.THEME['border_header']),
            right=Side(style='thin', color=self.THEME['border_header']),
            top=Side(style='medium', color=self.THEME['border_header']),
            bottom=Side(style='medium', color=self.THEME['border_header'])
        )
        
        # 内容单元格样式
        self.cell_font = Font(
            name='微软雅黑',
            color=self.THEME['dark_text'],
            size=10
        )
        self.cell_alignment = Alignment(
            horizontal="left",
            vertical="top",
            wrap_text=True,
            indent=1
        )
        self.cell_alignment_center = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )
        self.thin_border = Border(
            left=Side(style='thin', color=self.THEME['border']),
            right=Side(style='thin', color=self.THEME['border']),
            top=Side(style='thin', color=self.THEME['border']),
            bottom=Side(style='thin', color=self.THEME['border'])
        )
        
        # 偶数行填充
        self.even_row_fill = PatternFill(
            start_color=self.THEME['row_even'],
            end_color=self.THEME['row_even'],
            fill_type="solid"
        )
    
    def export_ui_test_cases(
        self,
        test_cases: List[TestCase],
        output_path: str,
        module_name: str = "TestCases"
    ) -> str:
        """
        导出 UI 测试用例
        
        Args:
            test_cases: 测试用例列表
            output_path: 输出路径
            module_name: 模块名称
        
        Returns:
            生成的文件路径
        """
        # 过滤 UI 相关的测试用例
        ui_cases = [
            tc for tc in test_cases
            if tc.type in [TestType.FUNCTIONAL, TestType.UI, TestType.SMOKE]
        ]
        
        return self._export(
            ui_cases, 
            output_path, 
            self.UI_COLUMNS, 
            f"{module_name}_功能测试用例",
            module_name, 
            include_locators=True
        )
    
    def export_all_test_cases(
        self,
        test_cases: List[TestCase],
        output_path: str,
        module_name: str = "TestCases"
    ) -> str:
        """
        导出全量测试用例
        
        Args:
            test_cases: 测试用例列表
            output_path: 输出路径
            module_name: 模块名称
        
        Returns:
            生成的文件路径
        """
        return self._export(
            test_cases, 
            output_path, 
            self.ALL_COLUMNS, 
            f"{module_name}_全量测试用例",
            module_name, 
            include_locators=False
        )
    
    def _export(
        self,
        test_cases: List[TestCase],
        output_path: str,
        columns: List[Dict],
        sheet_name: str,
        module_name: str,
        include_locators: bool
    ) -> str:
        """内部导出方法"""
        wb = openpyxl.Workbook()
        
        # 创建测试用例 Sheet
        ws = wb.active
        ws.title = sheet_name[:31]  # Excel sheet 名称最长31字符
        
        # 设置打印和显示选项
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        
        # 写入表头
        self._write_headers(ws, columns)
        
        # 写入数据
        for row_idx, tc in enumerate(test_cases, start=2):
            self._write_test_case_row(ws, row_idx, tc, columns, include_locators)
        
        # 设置列宽
        for col_idx, col in enumerate(columns, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = col['width']
        
        # 冻结表头
        ws.freeze_panes = 'A2'
        
        # 设置行高（表头）
        ws.row_dimensions[1].height = 35
        
        # 创建统计摘要 Sheet
        summary_ws = wb.create_sheet("统计摘要")
        self._write_summary(summary_ws, test_cases, module_name)
        
        # 保存文件
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(path))
        
        return str(path)
    
    def _write_headers(self, ws, columns: List[Dict]):
        """写入表头"""
        for col_idx, col in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col['name'])
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.header_alignment
            cell.border = self.header_border
    
    def _write_test_case_row(
        self,
        ws,
        row_idx: int,
        tc: TestCase,
        columns: List[Dict],
        include_locators: bool
    ):
        """写入测试用例行"""
        tc_dict = tc.to_dict()
        is_even_row = row_idx % 2 == 0
        
        # 格式化步骤 - 每步换行并包含定位器信息
        steps_text = self._format_steps_with_locators(tc_dict, include_locators)
        
        # 构建行数据
        row_data = {
            "id": tc_dict['id'],
            "module": tc_dict['module'],
            "submodule": self._extract_submodule(tc_dict['title']),
            "title": tc_dict['title'],
            "steps": steps_text,
            "expected_result": tc_dict['expected_result'],
            "priority": tc_dict['priority'],
            "type": self._get_type_display_name(tc_dict['type']),
            "test_data": self._format_test_data(tc_dict['test_data']),
            "notes": "",
        }
        
        # 计算行高 - 根据步骤数量动态调整
        step_count = len(tc_dict.get('steps_detail', []))
        row_height = max(30, step_count * 25 + 10)
        ws.row_dimensions[row_idx].height = row_height
        
        for col_idx, col in enumerate(columns, start=1):
            key = col['key']
            value = row_data.get(key, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            
            # 基础样式
            cell.font = self.cell_font
            cell.border = self.thin_border
            
            # 交替行颜色
            if is_even_row:
                cell.fill = self.even_row_fill
            
            # 特殊列样式处理
            if key == "priority":
                self._apply_priority_style(cell, value)
            elif key == "type":
                self._apply_type_style(cell, value)
            elif key in ["id", "module", "submodule"]:
                cell.alignment = self.cell_alignment_center
            elif key == "steps":
                # 步骤列使用特殊对齐
                cell.alignment = Alignment(
                    horizontal="left",
                    vertical="top",
                    wrap_text=True,
                    indent=0
                )
            else:
                cell.alignment = self.cell_alignment
    
    def _format_steps_with_locators(self, tc_dict: Dict, include_locators: bool) -> str:
        """
        格式化测试步骤，包含定位器信息
        
        每个步骤格式:
        1. 操作描述
           [定位器: xxx]  (如果有定位器)
           → 预期结果
        """
        steps_detail = tc_dict.get('steps_detail', [])
        if not steps_detail:
            # 兼容旧格式
            return '\n'.join(tc_dict.get('steps', []))
        
        formatted_lines = []
        
        for step in steps_detail:
            order = step.get('order', 0)
            action = step.get('action', '')
            expected = step.get('expected', '')
            locator = step.get('locator', '')
            data = step.get('data', '')
            
            # 步骤操作
            step_line = f"{order}. {action}"
            formatted_lines.append(step_line)
            
            # 定位器信息（如果有且需要显示）
            if include_locators and locator:
                locator_line = f"   📍 定位器: {locator}"
                formatted_lines.append(locator_line)
            
            # 测试数据（如果有）
            if data:
                data_line = f"   📝 数据: {data}"
                formatted_lines.append(data_line)
            
            # 预期结果
            if expected:
                expected_line = f"   → 预期: {expected}"
                formatted_lines.append(expected_line)
            
            # 步骤间空行
            formatted_lines.append("")
        
        # 移除最后的空行
        if formatted_lines and formatted_lines[-1] == "":
            formatted_lines.pop()
        
        return '\n'.join(formatted_lines)
    
    def _extract_submodule(self, title: str) -> str:
        """从标题提取子模块（如果有）"""
        if '-' in title:
            parts = title.split('-', 1)
            return parts[0].strip()
        return ""
    
    def _get_type_display_name(self, type_value: str) -> str:
        """获取测试类型的显示名称"""
        type_names = {
            "functional": "功能测试",
            "ui": "UI测试",
            "smoke": "冒烟测试",
            "boundary": "边界测试",
            "negative": "异常测试",
            "security": "安全测试",
            "performance": "性能测试",
            "regression": "回归测试",
        }
        return type_names.get(type_value, type_value)
    
    def _apply_priority_style(self, cell, priority: str):
        """应用优先级样式"""
        colors = self.PRIORITY_COLORS.get(priority)
        if colors:
            cell.fill = PatternFill(
                start_color=colors['bg'],
                end_color=colors['bg'],
                fill_type="solid"
            )
            cell.font = Font(
                name='微软雅黑',
                bold=True,
                color=colors['text'],
                size=10
            )
            cell.alignment = self.cell_alignment_center
            cell.border = Border(
                left=Side(style='thin', color=colors['border']),
                right=Side(style='thin', color=colors['border']),
                top=Side(style='thin', color=colors['border']),
                bottom=Side(style='thin', color=colors['border'])
            )
    
    def _apply_type_style(self, cell, type_name: str):
        """应用测试类型样式"""
        colors = self.TYPE_COLORS.get(type_name)
        if colors:
            cell.fill = PatternFill(
                start_color=colors['bg'],
                end_color=colors['bg'],
                fill_type="solid"
            )
            cell.font = Font(
                name='微软雅黑',
                color=colors['text'],
                size=10
            )
        cell.alignment = self.cell_alignment_center
    
    def _format_test_data(self, test_data: Dict) -> str:
        """格式化测试数据"""
        if not test_data:
            return ""
        
        lines = []
        for key, value in test_data.items():
            lines.append(f"• {key}: {value}")
        return '\n'.join(lines)
    
    def _write_summary(self, ws, test_cases: List[TestCase], module_name: str):
        """写入统计摘要"""
        # ===== 标题区域 =====
        title_font = Font(name='微软雅黑', bold=True, size=16, color=self.THEME['primary'])
        subtitle_font = Font(name='微软雅黑', size=11, color=self.THEME['light_text'])
        section_font = Font(name='微软雅黑', bold=True, size=12, color=self.THEME['primary'])
        
        ws['A1'] = f"📊 {module_name} - 测试用例统计摘要"
        ws['A1'].font = title_font
        ws.merge_cells('A1:D1')
        ws.row_dimensions[1].height = 35
        
        ws['A2'] = f"📅 生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        ws['A2'].font = subtitle_font
        
        ws['A3'] = f"📋 总用例数: {len(test_cases)}"
        ws['A3'].font = subtitle_font
        
        # ===== 按优先级统计 =====
        ws['A5'] = "🎯 按优先级分布"
        ws['A5'].font = section_font
        
        # 表头
        headers = ['优先级', '数量', '占比', '说明']
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=6, column=col, value=header)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.header_alignment
            cell.border = self.header_border
        
        priority_desc = {
            'P0': '阻塞级 - 核心功能',
            'P1': '高优先级 - 重要功能',
            'P2': '中优先级 - 一般功能',
            'P3': '低优先级 - 边缘功能'
        }
        
        row = 7
        total = len(test_cases)
        for priority in Priority:
            count = len([tc for tc in test_cases if tc.priority == priority])
            pct = f"{count/total*100:.1f}%" if total > 0 else "0%"
            
            ws.cell(row=row, column=1, value=priority.value)
            ws.cell(row=row, column=2, value=count)
            ws.cell(row=row, column=3, value=pct)
            ws.cell(row=row, column=4, value=priority_desc.get(priority.value, ''))
            
            # 应用样式
            for col in range(1, 5):
                cell = ws.cell(row=row, column=col)
                cell.font = self.cell_font
                cell.border = self.thin_border
                cell.alignment = self.cell_alignment_center
                
                if col == 1:
                    self._apply_priority_style(cell, priority.value)
            
            row += 1
        
        # ===== 按类型统计 =====
        ws.cell(row=row + 1, column=1, value="📑 按测试类型分布")
        ws.cell(row=row + 1, column=1).font = section_font
        
        row += 2
        # 表头
        type_headers = ['测试类型', '数量', '占比']
        for col, header in enumerate(type_headers, start=1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.header_alignment
            cell.border = self.header_border
        
        row += 1
        for test_type in TestType:
            count = len([tc for tc in test_cases if tc.type == test_type])
            if count > 0:
                pct = f"{count/total*100:.1f}%" if total > 0 else "0%"
                type_name = self._get_type_display_name(test_type.value)
                
                ws.cell(row=row, column=1, value=type_name)
                ws.cell(row=row, column=2, value=count)
                ws.cell(row=row, column=3, value=pct)
                
                for col in range(1, 4):
                    cell = ws.cell(row=row, column=col)
                    cell.font = self.cell_font
                    cell.border = self.thin_border
                    cell.alignment = self.cell_alignment_center
                    
                    if col == 1:
                        self._apply_type_style(cell, type_name)
                
                row += 1
        
        # 设置列宽
        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 10
        ws.column_dimensions['C'].width = 10
        ws.column_dimensions['D'].width = 25
    
    def export(self, testcases: List[Dict], output_path: str) -> str:
        """
        通用导出方法 - 直接从字典列表导出
        
        Args:
            testcases: 测试用例字典列表
            output_path: 输出路径
        
        Returns:
            生成的文件路径
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "测试用例"
        
        if not testcases:
            ws['A1'] = "暂无测试用例数据"
            path = Path(output_path)
            path.parent.mkdir(parents=True, exist_ok=True)
            wb.save(str(path))
            return str(path)
        
        # 使用UI列配置
        columns = self.UI_COLUMNS
        
        # 写入表头
        self._write_headers(ws, columns)
        
        # 写入数据
        for row_idx, tc_dict in enumerate(testcases, start=2):
            is_even_row = row_idx % 2 == 0
            
            # 格式化步骤
            steps_text = self._format_steps_with_locators(tc_dict, include_locators=True)
            
            # 构建行数据
            row_data = {
                "id": tc_dict.get('id', ''),
                "module": tc_dict.get('module', ''),
                "submodule": self._extract_submodule(tc_dict.get('title', '')),
                "title": tc_dict.get('title', ''),
                "steps": steps_text,
                "expected_result": tc_dict.get('expected_result', ''),
                "priority": tc_dict.get('priority', 'P1'),
                "type": self._get_type_display_name(tc_dict.get('type', 'functional')),
                "test_data": self._format_test_data(tc_dict.get('test_data', {})),
                "notes": tc_dict.get('notes', ''),
            }
            
            # 计算行高
            step_count = len(tc_dict.get('steps_detail', tc_dict.get('steps', [])))
            row_height = max(30, step_count * 25 + 10)
            ws.row_dimensions[row_idx].height = row_height
            
            for col_idx, col in enumerate(columns, start=1):
                key = col['key']
                value = row_data.get(key, "")
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                
                # 基础样式
                cell.font = self.cell_font
                cell.border = self.thin_border
                
                if is_even_row:
                    cell.fill = self.even_row_fill
                
                # 特殊列样式
                if key == "priority":
                    self._apply_priority_style(cell, value)
                elif key == "type":
                    self._apply_type_style(cell, value)
                elif key in ["id", "module", "submodule"]:
                    cell.alignment = self.cell_alignment_center
                elif key == "steps":
                    cell.alignment = Alignment(
                        horizontal="left",
                        vertical="top",
                        wrap_text=True,
                        indent=0
                    )
                else:
                    cell.alignment = self.cell_alignment
        
        # 设置列宽
        for col_idx, col in enumerate(columns, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = col['width']
        
        # 冻结表头
        ws.freeze_panes = 'A2'
        ws.row_dimensions[1].height = 35
        
        # 保存
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(path))
        
        return str(path)


# ==================== 便捷函数 ====================

def export_ui_test_cases(
    test_cases: List[TestCase],
    output_dir: str,
    module_name: str
) -> str:
    """
    导出 UI 测试用例到 Excel
    
    Args:
        test_cases: 测试用例列表
        output_dir: 输出目录
        module_name: 模块名称
    
    Returns:
        生成的文件路径
    """
    exporter = ExcelExporter()
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    filename = f"{module_name}_UI_TestCases_{timestamp}.xlsx"
    output_path = Path(output_dir) / filename
    
    return exporter.export_ui_test_cases(test_cases, str(output_path), module_name)


def export_all_test_cases(
    test_cases: List[TestCase],
    output_dir: str,
    module_name: str
) -> str:
    """
    导出全量测试用例到 Excel
    
    Args:
        test_cases: 测试用例列表
        output_dir: 输出目录
        module_name: 模块名称
    
    Returns:
        生成的文件路径
    """
    exporter = ExcelExporter()
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    filename = f"{module_name}_All_TestCases_{timestamp}.xlsx"
    output_path = Path(output_dir) / filename
    
    return exporter.export_all_test_cases(test_cases, str(output_path), module_name)
